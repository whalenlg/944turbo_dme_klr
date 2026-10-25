"""Build a KLR disassembly workbook for another KLR image from 87KLR951.xlsx.

usage: build_klr_variant.py <src.xlsx> <new .mem> <reference .mem the src sheet describes> <out.xlsx> [new tag] [ref tag]
  e.g. build_klr_variant.py 87KLR951.xlsx 89KLR_951.mem 87KLR_951.mem 89KLR951.xlsx 89 87
  (dme_klr/asm_select.sh runs this automatically when the source workbook changes)

The KLR sheet covers the code only, with gaps for tables and unused pages,
so this works on its contiguous stretches of rows:
* A stretch whose bytes are the same in both images is copied as-is.
* Inside a stretch with changes, rows whose bytes match are copied with every
  annotation, and the rest is re-disassembled from the new bytes and
  highlighted. Where the new code is the same instruction sequence as the
  reference but moved, the label / description follow the code.
* Differences in the gaps (maps, tables) are listed on the diff sheet.
"""
import difflib
import re
import sys
from collections import Counter

import openpyxl
from openpyxl.styles import PatternFill, Font

from dis8048 import decode, fmt_ops

SRC, MEMNEW, MEMREF, OUT = sys.argv[1:5]
NEW = sys.argv[5] if len(sys.argv) > 5 else '89'
REF = sys.argv[6] if len(sys.argv) > 6 else '87'
LABPFX = 'label%s_' % NEW[-2:]
SIZE = 4096


def mem(path):
    out = bytearray(b'\xff' * SIZE)
    a = 0
    for tok in open(path).read().split():
        if tok.startswith('@'):
            a = int(tok[1:], 16)
        else:
            out[a] = int(tok, 16)
            a += 1
    return bytes(out)


bnew, bref = mem(MEMNEW), mem(MEMREF)

wb = openpyxl.load_workbook(SRC)
ws = wb['ASM']
NCOL = ws.max_column            # Addr .. label_rationale
DIFFCOL = NCOL + 1
ANNOT = range(6, NCOL)          # long_description, label_rationale (0-based)


def op_bytes(v):
    if isinstance(v, (int, float)):
        v = '%02d' % v
    toks = str(v or '').split()
    out = []
    for t in toks:
        try:
            out.append(int(t, 16))
        except ValueError:
            out.append(None)   # hand-typed value such as "OE"
    return out


# ---- source rows, with the address each one really covers ----------------------
src = []   # dict(r, v, a, ln)
for r in range(2, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, NCOL + 1)]
    if v[0] is None:
        continue
    a = int(str(v[0]).strip().rstrip(':'), 16)
    ln = max(1, len(op_bytes(v[2])))
    # A row that repeats the previous address (a typo in the sheet) really
    # follows the previous instruction
    if src and a == src[-1]['a']:
        a = src[-1]['a'] + src[-1]['ln']
    src.append(dict(r=r, v=v, a=a, ln=ln))
label_at = {s['a']: str(s['v'][1]).strip().rstrip(':') for s in src if s['v'][1]}

# contiguous stretches of rows
runs = []
for s in src:
    if runs and s['a'] == runs[-1][-1]['a'] + runs[-1][-1]['ln']:
        runs[-1].append(s)
    else:
        runs.append([s])


def key(mn, args):
    return (mn,) + tuple(a if a[0] != 'code' else ('code',) for a in args)


# ---- walk the stretches ---------------------------------------------------------
out = []   # dict(a, v | None, ins, status, src_row)
for run in runs:
    lo, hi = run[0]['a'], run[-1]['a'] + run[-1]['ln']
    if bnew[lo:hi] == bref[lo:hi]:
        out += [dict(a=s['a'], v=list(s['v']), status=None, ins=decode(bnew, s['a'])) for s in run]
        continue
    by_a = {s['a']: s for s in run}
    # Code that grew pushes the stretch's last instructions past its old end:
    # follow the reference's final bytes to where they now sit, up to the
    # next stretch
    nxt = next((r[0]['a'] for r in runs if r[0]['a'] >= hi), SIZE)
    tail = bref[max(lo, hi - 6):hi]
    p = bnew.find(tail, max(lo, hi - 6), min(nxt, hi + 64))
    if p > hi - len(tail):
        hi = p + len(tail)
    pc, part = lo, []
    while pc < hi:
        s = by_a.get(pc)
        if s and bnew[pc:pc + s['ln']] == bref[pc:pc + s['ln']]:
            part.append(dict(a=pc, v=list(s['v']), status=None, ins=decode(bnew, pc)))
            pc += s['ln']
            continue
        mn, ln, b, args = decode(bnew, pc)
        if pc + ln > hi:
            mn, ln, b, args = '.db', 1, [bnew[pc]], [('reg', '0x%02x' % bnew[pc])]
        part.append(dict(a=pc, v=None, status='changed', ins=(mn, ln, b, args)))
        pc += ln
    # carry labels/descriptions to code that moved within this stretch
    ref_seq = [(s['a'], key(*[decode(bref, s['a'])[i] for i in (0, 3)])) for s in run]
    new_seq = [(p['a'], key(p['ins'][0], p['ins'][3])) for p in part]
    sm = difflib.SequenceMatcher(None, [k for _, k in ref_seq], [k for _, k in new_seq], autojunk=False)
    for blk in sm.get_matching_blocks():
        if blk.size < 3:
            continue
        for j in range(blk.size):
            p = part[blk.b + j]
            a_ref = ref_seq[blk.a + j][0]
            if p['status'] and a_ref != p['a']:
                p['status'], p['from'] = 'moved', a_ref
    for p in part:   # same opcode and length at the same address, operand changed
        s = by_a.get(p['a'])
        if p['status'] == 'changed' and s and bnew[p['a']] == bref[p['a']] and s['ln'] == p['ins'][1]:
            p['status'], p['from'] = 'patched', p['a']
    out += part

by_a_src = {}
for s in src:
    by_a_src.setdefault(s['a'], s)
starts = {o['a'] for o in out}

# ---- labels ---------------------------------------------------------------------
lab, used = {}, set()
for o in out:
    if o['status'] is None and o['a'] in label_at:
        lab[o['a']] = label_at[o['a']]; used.add(label_at[o['a']])
for o in out:
    if o['status'] in ('moved', 'patched'):
        n = label_at.get(o['from'])
        if n and n not in used:
            lab[o['a']] = n; used.add(n)


def target_name(t):
    if t not in lab:
        if t not in starts:
            return 'label%04X' % t
        lab[t] = LABPFX + '%04X' % t
    return lab[t]


stats = Counter()
for o in out:
    mn, ln, b, args = o['ins']
    if o['status'] is None:
        v = o['v']
        for arg in args:   # copied code whose target got a new label
            if arg[0] == 'code' and arg[1] in label_at and target_name(arg[1]) != label_at[arg[1]]:
                for ci in (4, 5):
                    if isinstance(v[ci], str):
                        v[ci] = re.sub(r'(?<![\w])%s(?![\w])' % re.escape(label_at[arg[1]]), target_name(arg[1]), v[ci])
        continue
    v = [None] * NCOL
    v[0] = '0x%04X' % o['a']
    v[2] = ' '.join('%02X' % x for x in b)
    v[3] = mn
    ops = fmt_ops(args, target_name)
    v[4] = ops or None
    v[5] = ('%s %s' % (mn, ops)).strip()
    note = []
    if o['status'] in ('moved', 'patched'):
        src_v = by_a_src[o['from']]['v']
        for ci in ANNOT:
            v[ci] = src_v[ci]
        note.append(('same code as %s 0x%04X (%s %s)' if o['status'] == 'moved'
                     else 'same instruction as %s, operand changed: was 0x%04X %s %s')
                    % (REF, o['from'], src_v[3] or '', src_v[4] or ''))
    s = by_a_src.get(o['a'])
    if s and o['status'] != 'patched':
        note.append('%s @ this addr: %s | %s %s' % (REF, s['v'][2], s['v'][3] or '', s['v'][4] or ''))
        if s['v'][6] and o['status'] != 'moved':
            note.append('%s desc: %s' % (REF, s['v'][6]))
    elif not s:
        note.append('%s: no instruction starts at this addr' % REF)
    o['v'], o['note'] = v, ' ; '.join(note)
for o in out:
    if o['a'] in lab and not o['v'][1]:
        o['v'][1] = lab[o['a']] + ':'

# ---- write ASM ------------------------------------------------------------------
fill_changed = PatternFill('solid', fgColor='FFFFF2CC')   # yellow: new/different code
fill_moved = PatternFill('solid', fgColor='FFDDEBF7')     # blue: same code as reference, moved
ws.delete_rows(2, ws.max_row - 1)
ws.cell(1, DIFFCOL, 'diff_vs_%s' % REF)
for r, o in enumerate(out, 2):
    for ci, val in enumerate(o['v'], 1):
        if val is not None:
            ws.cell(r, ci, val)
    stats[o['status']] += 1
    if o['status']:
        ws.cell(r, DIFFCOL, o['note'])
        f = fill_moved if o['status'] == 'moved' else fill_changed
        for ci in range(1, DIFFCOL + 1):
            ws.cell(r, ci).fill = f
ws.column_dimensions[openpyxl.utils.get_column_letter(DIFFCOL)].width = 70

# ---- diff sheet -----------------------------------------------------------------
covered = set()
for s in src:
    covered.update(range(s['a'], s['a'] + s['ln']))
named = sorted((a, l) for a, l in label_at.items())
runs_d, s0 = [], None
for k in range(SIZE + 1):
    d = k < SIZE and bnew[k] != bref[k]
    if d and s0 is None:
        s0 = k
    if not d and s0 is not None:
        runs_d.append([s0, k]); s0 = None
merged = []
for a, e in runs_d:
    if merged and a - merged[-1][1] <= 4:
        merged[-1][1] = e
    else:
        merged.append([a, e])
wd = wb.create_sheet(('diff_%s_vs_%s' % (NEW, REF))[:31], 1)
wd.append(['start', 'end', 'len', 'in ASM sheet', 'nearest label at/before start (%s addr)' % REF,
           '%s bytes' % REF, '%s bytes' % NEW])
for c in wd[1]:
    c.font = Font(bold=True)
for a, e in merged:
    near = [n for n in named if n[0] <= a]
    inside = sum(1 for i in range(a, e) if i in covered)
    wd.append(['0x%04X' % a, '0x%04X' % (e - 1), e - a,
               'yes' if inside == e - a else ('no (table/data)' if not inside else '%d of %d bytes' % (inside, e - a)),
               '%s (0x%04X)' % (near[-1][1], near[-1][0]) if near else '',
               ' '.join('%02X' % x for x in bref[a:e]), ' '.join('%02X' % x for x in bnew[a:e])])
wd.append([])
wd.append(['%s reference image: %s (the bytes %s describes)' % (REF, MEMREF.split('/')[-1], SRC.split('/')[-1])])
wd.append(['%s image: %s' % (NEW, MEMNEW.split('/')[-1])])
wd.append(['ASM colours: yellow = new/different in %s, blue = same code as %s but at a new address (labels/comments carried over)' % (NEW, REF)])
for col, w in zip('ABCDEFG', (9, 9, 6, 16, 40, 60, 60)):
    wd.column_dimensions[col].width = w

wb.save(OUT)
print('rows', len(out), dict(stats), 'diff regions', len(merged),
      'new labels', sum(1 for l in lab.values() if l.startswith(LABPFX)))
