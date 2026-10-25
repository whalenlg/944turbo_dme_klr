"""Build a disassembly workbook for another DME image from disassembleperformance.xlsx.

usage: build_variant.py <src.xlsx> <new .mem> <reference .bin the src sheet was built from> <out.xlsx> [new tag] [ref tag]
  e.g. build_variant.py disassembleperformance.xlsx 28PIN_86DME.mem 28PIN_DME_PERFORMANCE.bin disassemble1986.xlsx 1986 perf
  (dme_klr/asm_select.sh runs this automatically when the source workbook changes)

* Rows whose address and bytes are identical in both images are copied from the
  source sheet with every annotation.
* Everything else is re-disassembled from the new bytes and highlighted.
* Where the new code is the same instruction sequence as the reference but moved, the
  reference label / long_description / register columns follow the code.
"""
import difflib
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import PatternFill, Font

from dis8051 import decode, fmt_numeric, hexh

SRC, BIN86, BINREF, OUT = sys.argv[1:5]
NEW = sys.argv[5] if len(sys.argv) > 5 else '1986'   # tag for the new image
REF = sys.argv[6] if len(sys.argv) > 6 else '1989'   # tag for the image the source sheet describes
LABPFX = 'label%s_' % NEW[-2:]

b86 = bytes(int(l, 16) for l in open(BIN86) if l.strip())
ref = open(BINREF, 'rb').read()
assert len(b86) == len(ref) == 8192

wb = openpyxl.load_workbook(SRC)
ws = wb['ASM']
NCOL = ws.max_column          # A..AH
DIFFCOL = NCOL + 1
ANNOT = range(12, NCOL)       # long_description .. bit4 (0-based)
REGIONS = ((0x0000, 0x1000), (0x1000, 0x2000))


def addr_of(v):
    return int(str(v).strip().rstrip(':'), 16)


# ---- source sheet rows -------------------------------------------------------
sheet = {}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value is not None:
        sheet[addr_of(ws.cell(r, 1).value)] = (r, [ws.cell(r, c).value for c in range(1, NCOL + 1)])
saddrs = sorted(sheet)
slen = {}
for i, a in enumerate(saddrs):
    nxt = saddrs[i + 1] if i + 1 < len(saddrs) else 0x2000
    slen[a] = nxt - a
label_at = {a: str(v[1]).strip().rstrip(':') for a, (_, v) in sheet.items() if v[1]}


def split_ops(s):
    return [t.strip() for t in str(s).split(',')] if s not in (None, '') else []


# ---- learn raw/mapped operand names from the source sheet ---------------------
raw_by_val = defaultdict(Counter)
mapped_by_raw = defaultdict(Counter)
for a in saddrs:
    v = sheet[a][1]
    mn, ln, b, args = decode(ref, a)
    if ln != slen[a] or v[3] != mn:
        continue
    raw, mp = split_ops(v[4]), split_ops(v[5])
    if len(raw) != len(args) or len(mp) != len(args):
        continue
    for arg, rt, mt in zip(args, raw, mp):
        k = 'bit' if arg[0] == 'nbit' else arg[0]
        if k in ('dir', 'bit'):
            rt, mt = rt.lstrip('/'), mt.lstrip('/')
            raw_by_val[(k, arg[1])][rt] += 1
            mapped_by_raw[(k, rt)][mt] += 1


def raw_name(kind, val):
    for name, _ in raw_by_val.get((kind, val), Counter()).most_common():
        # skip one-off hand spellings like 095h / 1AH
        if re.fullmatch(r'[0-9A-F]{2}h', name) or not re.fullmatch(r'[0-9A-Fa-f]+[hH]', name):
            return name
    return hexh(val)


def mapped_name(kind, raw):
    c = mapped_by_raw.get((kind, raw))
    return c.most_common(1)[0][0] if c else raw


# ---- walk the new image ------------------------------------------------------
ins86 = []   # dict(pc, ln, mn, b, args, copied)
for lo, hi in REGIONS:
    pc = lo
    while pc < hi:
        if pc in sheet and b86[pc:pc + slen[pc]] == ref[pc:pc + slen[pc]]:
            ln = slen[pc]
            mn, dl, b, args = decode(b86, pc)
            ins86.append(dict(pc=pc, ln=ln, mn=mn, b=list(b86[pc:pc + ln]),
                              args=args if dl == ln else [], copied=True))
            pc += ln
            continue
        if pc == 0x1ffe:   # trailing checksum word, shown as END like the source sheet
            w = (b86[0x1ffe] << 8) | b86[0x1fff]
            ins86.append(dict(pc=pc, ln=2, mn='END', b=[b86[0x1ffe], b86[0x1fff]], args=[], copied=False, word=w))
            break
        mn, ln, b, args = decode(b86, pc)
        lim = 0x1ffe if hi == 0x2000 else hi
        if pc + ln > lim:
            mn, ln, b, args = 'db', lim - pc, list(b86[pc:lim]), []
        ins86.append(dict(pc=pc, ln=ln, mn=mn, b=list(b), args=args, copied=False))
        pc += ln
by_pc = {i['pc']: i for i in ins86}
assert all(i['copied'] for i in ins86 if i['pc'] < 0x1000), 'lower 4K differs'

# ---- align moved code (upper half) -------------------------------------------
def key(mn, args, raw):
    if mn in ('db', 'END'):
        return ('raw', tuple(raw))
    return (mn,) + tuple(a if a[0] != 'code' else ('code',) for a in args)


seq89 = []
for a in saddrs:
    if a < 0x1000:
        continue
    mn, ln, b, args = decode(ref, a)
    raw = tuple(ref[a:a + slen[a]])
    seq89.append((a, key(mn, args, raw) if ln == slen[a] else ('raw', raw)))
seq86 = [(i['pc'], key(i['mn'], i['args'], i['b'])) for i in ins86 if i['pc'] >= 0x1000]

sm = difflib.SequenceMatcher(None, [k for _, k in seq89], [k for _, k in seq86], autojunk=False)
moved = {}  # 86 pc -> 89 pc
for blk in sm.get_matching_blocks():
    if blk.size < 4:
        continue
    for k in range(blk.size):
        a89, a86 = seq89[blk.a + k][0], seq86[blk.b + k][0]
        if not by_pc[a86]['copied'] and a89 != a86:
            moved[a86] = a89

# same instruction slot (same addr, mnemonic and length) with different operands,
# e.g. a jump-table vector or an immediate constant that changed
patched = {}
for i in ins86:
    a = i['pc']
    if i['copied'] or a in moved or a not in sheet or a < 0x1000:
        continue
    mn89, ln89, _, _ = decode(ref, a)
    if b86[a] == ref[a] and ln89 == i['ln'] == slen[a] and sheet[a][1][3] == mn89:
        patched[a] = a

# ---- labels -------------------------------------------------------------------
lab86, used = {}, set()
for i in ins86:
    if i['copied'] and i['pc'] in label_at:
        lab86[i['pc']] = label_at[i['pc']]; used.add(label_at[i['pc']])
for a86, a89 in list(moved.items()) + list(patched.items()):
    n = label_at.get(a89)
    if n and n not in used:
        lab86[a86] = n; used.add(n)


def target_name(t, pc):
    if t == pc:
        return '$'
    if t not in lab86:
        if t not in by_pc:
            return '0x%04x' % t
        lab86[t] = LABPFX + '%04x' % t
    return lab86[t]


def fmt_raw(arg, pc):
    k, val = arg
    if k == 'reg':
        return val
    if k == 'dir':
        return raw_name('dir', val)
    if k == 'bit':
        return raw_name('bit', val)
    if k == 'nbit':
        return '/' + raw_name('bit', val)
    if k == 'imm':
        return '#' + hexh(val)
    if k == 'imm16':
        return '#' + hexh(val, 4)
    return target_name(val, pc)


def fmt_mapped(arg, raw):
    if arg[0] in ('dir', 'bit'):
        return mapped_name(arg[0], raw)
    if arg[0] == 'nbit':
        return '/' + mapped_name('bit', raw[1:])
    return raw


rows_out = []   # (pc, values, status, note)   status: None | 'changed' | 'moved'
for i in ins86:
    pc, mn, b, args = i['pc'], i['mn'], i['b'], i['args']
    if i['copied']:
        v = list(sheet[pc][1])
        for arg in args:   # code refs whose label changed in the new image
            if arg[0] == 'code' and arg[1] in label_at:
                old, new = label_at[arg[1]], target_name(arg[1], pc)
                if new != old:
                    for ci in (4, 5):
                        if isinstance(v[ci], str):
                            v[ci] = re.sub(r'(?<![\w])%s(?![\w])' % re.escape(old), new, v[ci])
        rows_out.append((pc, v, None, None))
        continue
    v = [None] * NCOL
    v[0] = '0x%04x:' % pc
    v[2] = ' '.join('%02x' % x for x in b)
    v[3] = v[9] = mn
    if mn == 'END':
        v[6] = '0x%04x' % i['word']
    elif args:
        raws = [fmt_raw(a, pc) for a in args]
        v[4] = v[10] = ', '.join(raws)
        v[5] = ', '.join(fmt_mapped(a, r) for a, r in zip(args, raws))
        nargs = [args[1], args[0]] if b[0] == 0x85 else args   # sheet shows 85 xx yy in byte order
        v[6] = ', '.join(fmt_numeric(a) for a in nargs)
    note = []
    status = 'changed'
    if pc in moved:
        a89 = moved[pc]
        o = sheet[a89][1]
        for ci in ANNOT:
            v[ci] = o[ci]
        status = 'moved'
        note.append('same code as %s 0x%04x (%s %s)' % (REF, a89, o[3] or '', o[4] or ''))
    elif pc in patched:
        o = sheet[pc][1]
        for ci in ANNOT:
            v[ci] = o[ci]
        note.append('same instruction as %s, operand changed: was %s | %s %s' % (REF, o[2], o[3] or '', o[4] or ''))
    if pc in sheet and pc not in patched:
        o = sheet[pc][1]
        note.append('%s @ this addr: %s | %s %s' % (REF, o[2], o[3] or '', o[4] or ''))
        if o[12] and status != 'moved':
            note.append('%s desc: %s' % (REF, o[12]))
    else:
        note.append('%s: no instruction starts at this addr' % REF)
    rows_out.append((pc, v, status, ' ; '.join(note)))

for pc, v, status, note in rows_out:
    if pc in lab86 and not v[1]:
        v[1] = lab86[pc] + ':'
    if status and v[1]:
        v[8] = v[1]

# ---- write ASM ----------------------------------------------------------------
fill_changed = PatternFill('solid', fgColor='FFFFF2CC')   # yellow: new/different code or data
fill_moved = PatternFill('solid', fgColor='FFDDEBF7')     # blue: same code as reference, moved
first_upper = min(r for a, (r, _) in sheet.items() if a >= 0x1000)
ws.delete_rows(first_upper, ws.max_row - first_upper + 1)
ws.cell(1, DIFFCOL, 'diff_vs_%s' % REF)
for pc, v, status, note in rows_out:      # lower half: rows stay in place
    if pc >= 0x1000:
        continue
    r = sheet[pc][0]
    for ci in (1, 4, 5):
        if ws.cell(r, ci + 1).value != v[ci]:
            ws.cell(r, ci + 1, v[ci])
r = first_upper
stats = Counter()
for pc, v, status, note in rows_out:
    if pc < 0x1000:
        continue
    for ci, val in enumerate(v, 1):
        if val is not None:
            ws.cell(r, ci, val)
    stats[status] += 1
    if status:
        ws.cell(r, DIFFCOL, note)
        f = fill_moved if status == 'moved' else fill_changed
        for ci in range(1, DIFFCOL + 1):
            ws.cell(r, ci).fill = f
    r += 1
ws.column_dimensions[openpyxl.utils.get_column_letter(DIFFCOL)].width = 70

# ---- Constants: new vs reference bytes --------------------------------------------
wc = wb['Constants']
cc = wc.max_column + 1
wc.cell(1, cc, '%s bytes' % NEW)
wc.cell(1, cc + 1, '%s bytes' % REF)
for rr in range(2, wc.max_row + 1):
    m = re.findall(r'([0-9A-Fa-f]{4})h', str(wc.cell(rr, 3).value or ''))
    if not m:
        continue
    lo, hi = int(m[0], 16), int(m[-1], 16)
    s86 = ' '.join('%02X' % x for x in b86[lo:hi + 1])
    s89 = ' '.join('%02X' % x for x in ref[lo:hi + 1])
    wc.cell(rr, cc, s86)
    wc.cell(rr, cc + 1, s89)
    if s86 != s89:
        wc.cell(rr, cc).fill = fill_changed
        wc.cell(rr, cc).font = Font(bold=True)

# ---- diff sheet ---------------------------------------------------------------
named = []
for sname, acol, ncol in (('Constants', 3, 4), ('Numbered Maps', 4, 2), ('Directly Addressed Tables', 1, 2)):
    s = wb[sname]
    for rr in range(2, s.max_row + 1):
        for m in re.findall(r'([0-9A-Fa-f]{4})h', str(s.cell(rr, acol).value or '')):
            named.append((int(m, 16), s.cell(rr, ncol).value, sname))
named += [(a, l, 'ASM label') for a, l in label_at.items() if a >= 0x1000]
named.sort()

runs, s0 = [], None
for k in range(8193):
    d = k < 8192 and b86[k] != ref[k]
    if d and s0 is None:
        s0 = k
    if not d and s0 is not None:
        runs.append([s0, k]); s0 = None
merged = []
for a, e in runs:
    if merged and a - merged[-1][1] <= 4:
        merged[-1][1] = e
    else:
        merged.append([a, e])

wd = wb.create_sheet(('diff_%s_vs_%s' % (NEW, REF))[:31], 1)
wd.append(['start', 'end', 'len', 'nearest named item at/before start (%s addr)' % REF, 'source', '%s bytes' % REF, '%s bytes' % NEW])
for c in wd[1]:
    c.font = Font(bold=True)
for a, e in merged:
    near = [n for n in named if n[0] <= a]
    n = near[-1] if near else None
    wd.append(['0x%04x' % a, '0x%04x' % (e - 1), e - a,
               '%s (0x%04x)' % (n[1], n[0]) if n else '', n[2] if n else '',
               ' '.join('%02X' % x for x in ref[a:e]), ' '.join('%02X' % x for x in b86[a:e])])
wd.append([])
wd.append(['%s reference image: %s (the bytes %s was built from)' % (REF, BINREF.split('/')[-1], SRC.split('/')[-1])])
wd.append(['%s image: %s' % (NEW, BIN86.split('/')[-1])])
wd.append(['ASM colours: yellow = new/different in %s, blue = same code as %s but at a new address (labels/comments carried over)' % (NEW, REF)])
for col, w in zip('ABCDEFG', (9, 9, 6, 46, 26, 60, 60)):
    wd.column_dimensions[col].width = w

wb.save(OUT)
print('rows', len(rows_out), dict(stats), 'diff regions', len(merged),
      'new labels', sum(1 for l in lab86.values() if l.startswith(LABPFX)),
      'moved labels', sum(1 for a in moved if a in lab86 and not lab86[a].startswith(LABPFX)))
